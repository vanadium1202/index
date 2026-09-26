"""Excel summary in the user's requested layout: one row per market, 7- and 15-day forecast
plus % accuracy (= 100% - backtest mean absolute % error). A Details sheet holds the numbers
the summary links to, so every summary cell is a formula."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Aptos Narrow"
NAMES = {"brent": "Brent", "wti": "WTI", "henry_hub": "Henry Hub"}
UNITS = {"brent": "$/bbl", "wti": "$/bbl", "henry_hub": "$/MMBtu"}
MODEL_LABEL = {"random_walk": "No-change (random walk)", "ridge": "Ridge regression",
               "gbm": "Gradient-boosted trees", "ensemble": "Ridge + GBM ensemble"}

BLUE = Font(name=FONT, size=12, color="0000FF")      # hardcoded model output
BLACK = Font(name=FONT, size=12)
GREEN = Font(name=FONT, size=12, color="008000")     # link to another sheet
BOLD = Font(name=FONT, size=12, bold=True)
NOTE = Font(name=FONT, size=10, italic=True, color="595959")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

DETAIL_COLS = [
    ("Market", 12), ("Unit", 9), ("Horizon (days)", 9), ("As of", 11), ("Target date", 11),
    ("Last price", 10), ("Forecast", 10), ("Change %", 9), ("80% low", 9), ("80% high", 9),
    ("95% low", 9), ("95% high", 9), ("P(rise)", 8), ("Backtest mean abs. error", 11),
    ("% accuracy", 10), ("Within ±5%", 9), ("80% range hit rate", 10), ("Model used", 22),
]


def write_excel_summary(forecasts: list[dict], path: Path, source: str) -> Path:
    fc = pd.DataFrame(forecasts)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    det = wb.create_sheet("Details")

    # ---- Details: one row per market x horizon (model outputs, blue) ----
    for j, (name, width) in enumerate(DETAIL_COLS, start=1):
        c = det.cell(row=1, column=j, value=name)
        c.font, c.fill, c.alignment, c.border = BOLD, HEAD_FILL, CENTER, BOX
        det.column_dimensions[c.column_letter].width = width
    det.row_dimensions[1].height = 45
    rows: dict[tuple[str, int], int] = {}
    for i, f in enumerate(fc.to_dict("records"), start=2):
        rows[(f["market"], int(f["horizon_calendar_days"]))] = i
        vals = [
            NAMES.get(f["market"], f["market"]), UNITS.get(f["market"], ""),
            int(f["horizon_calendar_days"]), pd.Timestamp(f["as_of"]).date(),
            pd.Timestamp(f["target_date"]).date(), f["last_price"], f["forecast_price"],
            f"=G{i}/F{i}-1", f["p10"], f["p90"], f["p2_5"], f["p97_5"], f["prob_up_%"] / 100,
            f["backtest_MAPE_%"] / 100, f"=1-N{i}", f["backtest_within_5%"] / 100,
            f["backtest_coverage80_%"] / 100, MODEL_LABEL.get(f["model"], f["model"]),
        ]
        fmts = ["@", "@", "0", "yyyy-mm-dd", "yyyy-mm-dd", "0.00", "0.00", "+0.0%;-0.0%;0.0%",
                "0.00", "0.00", "0.00", "0.00", "0%", "0.0%", "0.0%", "0%", "0%", "@"]
        for j, (v, fmt) in enumerate(zip(vals, fmts), start=1):
            c = det.cell(row=i, column=j, value=v)
            is_formula = isinstance(v, str) and v.startswith("=")
            c.font = BLACK if is_formula or j <= 5 else BLUE
            c.number_format, c.border = fmt, BOX
            c.alignment = Alignment(horizontal="left" if j in (1, 2, 18) else "center")
    last = 1 + len(fc)
    notes = [
        f"Data: {source}. Blue = model output (python run_forecast.py); black = formula.",
        "% accuracy = 1 - backtest mean absolute % error of the forecast price "
        "(8-year walk-forward, out-of-sample).",
        "80%/95% range: the actual price is expected inside it 80%/95% of the time; "
        "'80% range hit rate' is how often that happened in the backtest.",
        "7 and 15 calendar days = 5 and 11 trading days after the as-of date.",
    ]
    for k, text in enumerate(notes):
        det.cell(row=last + 2 + k, column=1, value=text).font = NOTE
    det.freeze_panes = "B2"

    # ---- Sheet1: the requested layout, every value linked to Details ----
    ws.column_dimensions["G"].width = 22
    for col in "HIJK":
        ws.column_dimensions[col].width = 13
    ws["H13"] = "Crude oil Price forecast per bbl"
    ws["H13"].font = Font(name=FONT, size=14, bold=True)
    heads = ["Parameter", "7 day", "% accuracy", "15 day", "% accuracy"]
    for j, h in enumerate(heads):
        c = ws.cell(row=15, column=7 + j, value=h)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, BOX
        c.alignment = CENTER if j else Alignment(horizontal="left", vertical="center")
    as_of = fc["as_of"].iloc[0]
    for i, market in enumerate(["brent", "wti", "henry_hub"], start=16):
        name = NAMES[market] + (" ($/MMBtu)" if market == "henry_hub" else "")
        c = ws.cell(row=i, column=7, value=name)
        c.font, c.border = BLACK, BOX
        for j, days in ((8, 7), (10, 15)):
            r = rows.get((market, days))
            price = ws.cell(row=i, column=j, value=f"=Details!G{r}" if r else "n/a")
            acc = ws.cell(row=i, column=j + 1, value=f"=Details!O{r}" if r else "n/a")
            for c, fmt in ((price, "0.00"), (acc, "0.0%")):
                c.font, c.number_format, c.border = GREEN, fmt, BOX
                c.alignment = Alignment(horizontal="center")
    sheet_notes = [
        f"Forecast made on prices up to {as_of}; 7 day = {fc.loc[fc.horizon_calendar_days == 7, 'target_date'].iloc[0]}, "
        f"15 day = {fc.loc[fc.horizon_calendar_days == 15, 'target_date'].iloc[0]}."
        if {7, 15} <= set(fc.horizon_calendar_days) else f"Forecast made on prices up to {as_of}.",
        "Prices in $/bbl except Henry Hub ($/MMBtu).",
        "% accuracy = 100% - average % error of the forecast in an 8-year out-of-sample backtest.",
        "Price ranges and probability of a rise: see the Details sheet.",
    ]
    for k, text in enumerate(sheet_notes):
        ws.cell(row=20 + k, column=7, value=text).font = NOTE
    ws.print_area = f"G13:K{20 + len(sheet_notes) - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    det.page_setup.orientation = "landscape"
    det.page_setup.fitToWidth = 1
    det.page_setup.fitToHeight = 0
    det.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return path
